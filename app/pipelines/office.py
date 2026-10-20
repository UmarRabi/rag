import io
from collections.abc import Iterable, Iterator
from typing import Any

import docx
import openpyxl
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

from app.domain.content import Extraction, RawContent, TextSegment
from app.domain.enums import MediaType
from app.ingestion.cleaning import clean_text
from app.pipelines.base import extraction

HEADING_STYLE_PREFIX = "Heading"
TITLE_STYLE = "Title"
DEFAULT_HEADING_LEVEL = 2
MAX_HEADING_LEVEL = 6
CELL_SEPARATOR = " | "


def heading_level(paragraph: Paragraph) -> int | None:
    name = (paragraph.style.name or "") if paragraph.style is not None else ""
    if name == TITLE_STYLE:
        return 1
    if not name.startswith(HEADING_STYLE_PREFIX) or not paragraph.text.strip():
        return None
    suffix = name.removeprefix(HEADING_STYLE_PREFIX).strip()
    return min(int(suffix), MAX_HEADING_LEVEL) if suffix.isdigit() else DEFAULT_HEADING_LEVEL


def segment(parts: Iterable[str], page: int | None = None) -> TextSegment | None:
    text = clean_text("\n\n".join(part for part in parts if part.strip()))
    return TextSegment(text, page=page) if text else None


def table_rows(rows: Iterable[Iterable[str]]) -> str:
    lines = (CELL_SEPARATOR.join(cell.strip() for cell in row if cell.strip()) for row in rows)
    return "\n".join(line for line in lines if line)


class DocxExtractor:
    media_type = MediaType.DOCX

    def extract(self, content: RawContent) -> Extraction:
        document = docx.Document(io.BytesIO(content.data))
        result = segment(self._blocks(document.iter_inner_content()))
        return extraction([result])

    def _blocks(self, items: Iterable[Paragraph | Table]) -> Iterator[str]:
        for item in items:
            if isinstance(item, Table):
                yield table_rows((cell.text for cell in row.cells) for row in item.rows)
            elif level := heading_level(item):
                yield f"{'#' * level} {item.text}"
            else:
                yield item.text


class PptxExtractor:
    media_type = MediaType.PPTX

    def extract(self, content: RawContent) -> Extraction:
        presentation = Presentation(io.BytesIO(content.data))
        segments = (
            segment(self._slide_parts(slide), page=number)
            for number, slide in enumerate(presentation.slides, start=1)
        )
        return extraction(segments)

    def _slide_parts(self, slide: Any) -> Iterator[str]:
        for shape in slide.shapes:
            yield from self._shape_text(shape)
        if slide.has_notes_slide:
            yield slide.notes_slide.notes_text_frame.text

    def _shape_text(self, shape: Any) -> Iterator[str]:
        if shape.has_text_frame:
            yield shape.text_frame.text
        if getattr(shape, "has_table", False):
            yield table_rows((cell.text for cell in row.cells) for row in shape.table.rows)


class XlsxExtractor:
    media_type = MediaType.XLSX

    def extract(self, content: RawContent) -> Extraction:
        workbook = openpyxl.load_workbook(io.BytesIO(content.data), read_only=True, data_only=True)
        try:
            segments = (
                segment(
                    [f"Sheet: {sheet.title}", *self._rows(sheet.iter_rows(values_only=True))],
                    number,
                )
                for number, sheet in enumerate(workbook.worksheets, start=1)
            )
            return extraction(segments)
        finally:
            workbook.close()

    def _rows(self, rows: Iterator[tuple[Any, ...]]) -> Iterator[str]:
        header = next(rows, None)
        if header is None:
            return
        names = [
            str(name) if name is not None else f"Column {index}"
            for index, name in enumerate(header, 1)
        ]
        for row in rows:
            cells = [
                f"{name}: {value}"
                for name, value in zip(names, row, strict=False)
                if value is not None
            ]
            if cells:
                yield CELL_SEPARATOR.join(cells)
