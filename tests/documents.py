import io
import wave
from typing import Any, Protocol

import av
import docx
import numpy
import openpyxl
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SAMPLE_WIDTH = 2


def pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def pdf_document(*pages: str) -> bytes:
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    kids = []
    for index, text in enumerate(pages):
        page_id, content_id = 4 + index * 2, 5 + index * 2
        stream = f"BT /F1 12 Tf 72 720 Td ({pdf_escape(text)}) Tj ET".encode() if text else b""
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode()
        kids.append(f"{page_id} 0 R")
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>".encode()

    output = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for number in sorted(objects):
        offsets[number] = len(output)
        output += f"{number} 0 obj\n".encode() + objects[number] + b"\nendobj\n"

    size = max(objects) + 1
    xref = len(output)
    output += f"xref\n0 {size}\n0000000000 65535 f \n".encode()
    for number in range(1, size):
        output += f"{offsets[number]:010d} 00000 n \n".encode()
    output += f"trailer\n<< /Size {size} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(output)


class Savable(Protocol):
    def save(self, target: Any, /) -> Any: ...


def saved(document: Savable) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def docx_document() -> bytes:
    document = docx.Document()
    document.add_heading("Returns policy", level=1)
    document.add_paragraph("Items can be returned within 30 days.")
    table = document.add_table(rows=2, cols=2)
    for row, (item, days) in enumerate([("Item", "Days"), ("Router", "14")]):
        table.cell(row, 0).text = item
        table.cell(row, 1).text = days
    return saved(document)


def pptx_document() -> bytes:
    presentation = Presentation()
    for title, body in [
        ("Onboarding", "Welcome to the team"),
        ("Benefits", "Health cover from day one"),
    ]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
        box.text_frame.text = body
    presentation.slides[1].notes_slide.notes_text_frame.text = "Mention the gym discount"
    return saved(presentation)


def xlsx_document() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Prices"
    for row in [("Product", "Price"), ("Router X", 120), ("Router Y", None)]:
        sheet.append(row)
    return saved(workbook)


def html_document() -> bytes:
    return b"""<!doctype html><html><head><title>Help</title></head><body>
    <nav><a href="/">Home</a> | <a href="/pricing">Pricing</a></nav>
    <article><h1>Resetting your router</h1>
    <p>Hold the reset button for ten seconds until the light blinks green.
    The router restarts and restores its factory settings.</p>
    <p>Reconnect using the network name printed on the label underneath.</p></article>
    <footer>Copyright Acme</footer></body></html>"""


def png_image(colour: str = "white", size: int = 64) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (size, size), colour).save(buffer, format="PNG")
    return buffer.getvalue()


def wav_audio(seconds: float = 1.0, rate: int = 16_000) -> bytes:
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as audio:
        audio.setnchannels(1)
        audio.setsampwidth(SAMPLE_WIDTH)
        audio.setframerate(rate)
        audio.writeframes(bytes(SAMPLE_WIDTH * int(seconds * rate)))
    return buffer.getvalue()


def mp4_video(
    colours: list[str], seconds_per_colour: float = 4.0, *, with_audio: bool = False
) -> bytes:
    buffer = io.BytesIO()
    fps = 4
    with av.open(buffer, mode="w", format="mp4") as container:
        video = container.add_stream("mpeg4", rate=fps)
        video.width, video.height, video.pix_fmt = 64, 64, "yuv420p"
        video.codec_context.gop_size = fps * 2
        audio = container.add_stream("aac", rate=16_000) if with_audio else None
        for index, colour in enumerate(colours):
            for step in range(int(seconds_per_colour * fps)):
                frame = av.VideoFrame.from_image(Image.new("RGB", (64, 64), colour))
                frame.pts = index * int(seconds_per_colour * fps) + step
                container.mux(video.encode(frame))
        container.mux(video.encode(None))
        if audio is not None:
            samples = int(len(colours) * seconds_per_colour * 16_000)
            silence = av.AudioFrame.from_ndarray(
                numpy.zeros((1, samples), dtype=numpy.float32),
                format="fltp",
                layout="mono",
            )
            silence.sample_rate = 16_000
            container.mux(audio.encode(silence))
            container.mux(audio.encode(None))
    return buffer.getvalue()
